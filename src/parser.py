import os 
from bs4 import BeautifulSoup
import pandas as pd
import numpy as np  
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill
# carpeta script es la ruta del script actual, y ".." significa “la carpeta de arriba”:

carpeta_script = os.path.dirname(os.path.abspath(__file__))
ruta_HTML= os.path.join(carpeta_script, "..", "data", "raw", "index.html")
hoy = datetime.now().strftime("%Y-%m-%d %H-%M-%S")
processed = os.path.join(carpeta_script, "..", "data", "processed")
os.makedirs(processed, exist_ok=True)
report = os.path.join(carpeta_script, "..", "data", "report")
os.makedirs(report, exist_ok=True)
libros = []
GUARDAR = True
PRUEBA = True

with open(ruta_HTML,  mode = 'r' , encoding="utf-8") as archivo:
    contenido = archivo.read()
    soup = BeautifulSoup(contenido, 'html.parser')
    for book in soup.find_all('article', class_='product_pod'):
        title = book.h3.a['title']
        price = book.find('p', class_='price_color').text
        libros.append({"title": title, "price": price})

df = pd.DataFrame(libros)
# print(df.head()) # vista previa de los primeros 5 registros
# print(df.dtypes) # muestra los tipos de datos de cada columna
df["price"] =  df["price"].str.replace('£', '').astype(float)
ruta_final = os.path.join(processed, f"precios{hoy}.csv")
# listar la carpeta
# leer csv anterior asi no se confunde por que ya los datos anteriores estan
# para luego guardar el actual

# Filtra solo los archivos que terminan en .csv
lista = []
for i in os.listdir(processed):
    if i.endswith(".csv"):
        lista.append(i)
if not lista:
    print("No hay archivos anteriores en la carpeta. Es la primera ejecución.")
    viejo = None
    anterior = None
else:
    reciente = sorted(lista)[-1]
    ruta  = os.path.join(processed, reciente)
    viejo = pd.read_csv(ruta)
    if viejo is not None:
        # df es el actual, viejo es el anterior
        r = pd.merge(df, viejo, how='outer', on="title", suffixes=('_actual', '_anterior'))
        if r is not None:
            r['price_anterior'] = r['price_anterior']
            r['price_actual'] = r['price_actual']
            r['price_cambio'] = r['price_actual'] - r['price_anterior']
            
            condicionales = [
            r['price_cambio'] > 0, 
            r['price_cambio'] < 0, 
            r['price_cambio'] == 0
            ]
            etiqutas = ['Subio', 'Bajo', 'Sin cambios']
            r['Estado'] = np.select(condicionales, etiqutas, default='Sin cambios')
            # print(r[r['Estado'] != 'Sin cambios']) # muestra los libros que cambiaron de precio
rp= r[r['Estado'] != 'Sin cambios']

if GUARDAR:
    df.to_csv(ruta_final, index=False) # guarda el DataFrame en un archivo CSV
if PRUEBA:
    ruta_report= os.path.join(report, "reporte_prueba.xlsx")
else:
    ruta_report= os.path.join(report, f"reporte{hoy}.xlsx")


# visual excel

if not rp.empty:
    rp.to_excel(ruta_report, index=False)
    wb = load_workbook(ruta_report)
    ws = wb.active
    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    negritas = Font(bold= True)
    ws['A1'].font = negritas
    ws['B1'].font = negritas
    ws['C1'].font = negritas
    ws['D1'].font = negritas
    ws['E1'].font = negritas
    relleno = PatternFill(patternType='solid', fgColor='00FF3E' )
    rellenoR = PatternFill(patternType='solid', fgColor='FF0100' )
    cells = ws['E']
    for cell in cells:
        E = cell.value
        if E == "Subio":
            cell.fill = relleno
        elif E == "Bajo":
            cell.fill = rellenoR
        else:
            ws['E']
    wb.save(ruta_report)
else:
    print(f"Hoy {hoy} no hubo cambios en los precios ")